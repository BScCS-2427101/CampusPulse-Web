import asyncio
import os
import random
import datetime
import shutil
import openpyxl
from pydantic import BaseModel
from typing import List, Dict, Optional
from app.services.cache import DataCache

cache = DataCache()

class LogEntry(BaseModel):
    timestamp: str
    dataset: str
    record_id: str
    field_changed: str
    old_value: str
    new_value: str

class LiveSimulationState:
    def __init__(self):
        self.is_running = False
        self.interval_ms = 10000
        self.update_count = 0
        self.last_update = "Never"
        self.log_entries: List[LogEntry] = []
        self._task: Optional[asyncio.Task] = None

state = LiveSimulationState()

def _get_live_path():
    # Adjust path assuming backend runs from backend dir and data is in ../../data/
    # The cache uses loader, let's just use the loader's live file path
    from app.services.data_loader import DataLoader
    loader = DataLoader()
    return loader.live_file

def _get_orig_path():
    from app.services.data_loader import DataLoader
    loader = DataLoader()
    return loader.original_file

def perform_one_update():
    file_path = _get_live_path()
    if not os.path.exists(file_path):
        raise Exception("Live dataset not found.")
        
    wb = openpyxl.load_workbook(file_path)
    
    targets = [
        ("Enrollments", ["assessment_score", "attendance_percentage", "assignment_completion_rate", "exam_attempts", "faculty_feedback_score"]),
        ("Learning_Activity", ["study_hours", "library_visits", "lms_logins", "video_minutes", "assignment_submissions", "quiz_attempts", "engagement_score"])
    ]
    
    sheet_name, fields = random.choice(targets)
    if sheet_name not in wb.sheetnames:
        return None
        
    sheet = wb[sheet_name]
    if sheet.max_row <= 1:
        return None
        
    headers = [cell.value for cell in sheet[1]]
    
    row_idx = random.randint(2, sheet.max_row)
    field_name = random.choice(fields)
    
    if field_name not in headers:
        return None
        
    col_idx = headers.index(field_name) + 1
    id_col_idx = 1 # Assuming ID is first column
    
    old_value = sheet.cell(row=row_idx, column=col_idx).value
    record_id = sheet.cell(row=row_idx, column=id_col_idx).value
    
    if old_value is None: old_value = 0
    
    # Bound and compute new value safely
    new_value = old_value
    if field_name in ["assessment_score", "attendance_percentage", "assignment_completion_rate"]:
        new_value = max(0, min(100, float(old_value) + random.randint(-5, 5)))
    elif field_name == "faculty_feedback_score":
        new_value = max(0.0, min(10.0, float(old_value) + random.uniform(-0.5, 0.5)))
        new_value = round(new_value, 1)
    else:
        # Count based
        new_value = max(0, int(old_value) + random.randint(-2, 3))
        
    sheet.cell(row=row_idx, column=col_idx, value=new_value)
    
    wb.save(file_path)
    wb.close()
    
    # Invalidate cache and reload
    cache.invalidate()
    cache.get_data() # Force rebuild
    
    state.update_count += 1
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    state.last_update = timestamp
    
    log_entry = LogEntry(
        timestamp=timestamp,
        dataset=sheet_name,
        record_id=str(record_id),
        field_changed=field_name,
        old_value=str(old_value),
        new_value=str(new_value)
    )
    
    state.log_entries.insert(0, log_entry)
    if len(state.log_entries) > 50:
        state.log_entries.pop()
        
    return log_entry

async def simulation_loop():
    while state.is_running:
        try:
            perform_one_update()
        except Exception as e:
            print(f"Simulation loop error: {e}")
            state.is_running = False
            break
        await asyncio.sleep(state.interval_ms / 1000.0)

def start_simulation(interval_ms: int):
    if state.is_running:
        return
    state.interval_ms = interval_ms
    state.is_running = True
    state._task = asyncio.create_task(simulation_loop())

def stop_simulation():
    state.is_running = False
    if state._task:
        state._task.cancel()
        state._task = None

def reset_dataset():
    stop_simulation()
    orig_path = _get_orig_path()
    live_path = _get_live_path()
    if not os.path.exists(orig_path):
        raise Exception("Original dataset missing. Cannot reset.")
        
    shutil.copy2(orig_path, live_path)
    cache.invalidate()
    cache.get_data() # Force rebuild
    
    state.update_count = 0
    state.last_update = "Reset"
    state.log_entries.clear()

def get_status():
    return {
        "is_running": state.is_running,
        "interval_ms": state.interval_ms,
        "update_count": state.update_count,
        "last_update": state.last_update,
        "log_entries": [e.dict() for e in state.log_entries],
        "cache_last_updated": cache.get_data()["last_updated"] if cache.last_updated else None
    }
